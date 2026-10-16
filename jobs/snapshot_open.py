import io
from datetime import date, timedelta, datetime
from collections import defaultdict
import openpyxl
from core import config, db
from core.iia_client import create_authenticated_session
from core.jobio import run_job, logger

# Daftar 24 cabang master
MASTER_KDCAB = {
    "G001", "G026", "G033", "G107", "G113", "G117", "G137", "G143",
    "G146", "G157", "G165", "G174", "G219", "G224", "G236", "G237",
    "G260", "G263", "G265", "G272", "G273", "G283", "G295", "G298"
}

def parse_date_val(value: any) -> date | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value if isinstance(value, date) else value.date()
    try:
        return datetime.strptime(str(value).strip(), "%d-%B-%Y").date()
    except Exception:
        return None

def download_rekap_ringkas(session) -> bytes:
    """Download laporan rekap intransit kumulatif (_isDetail=false) dari 1 Jan s/d H-1."""
    today = date.today()
    h1 = today - timedelta(days=1)
    start_date = date(today.year, 1, 1)

    start_str = f"{start_date.day}-{start_date.strftime('%B-%Y')}"
    end_str = f"{h1.day}-{h1.strftime('%B-%Y')}"

    url = (
        f"{config.IIA_BASE_URL}/reports/printRekapIntransit"
        f"?_startDate={start_str}&_endDate={end_str}"
        f"&_respId={config.IIA_RESP_ID}&_isDetail=false"
    )
    logger.info(f"Download rekap kumulatif ({start_str} s/d {end_str})...")
    resp = session.get(url, timeout=120)
    if resp.status_code != 200:
        raise ValueError(f"HTTP {resp.status_code} saat download rekap ringkas")
    logger.info(f"✓ Download rekap selesai ({len(resp.content)/1024:.1f} KB)")
    return resp.content

def parse_and_store_snapshot(content: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active

    header_row_idx = None
    col_headers = []
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=25, values_only=True), start=1):
        if row[0] == "NO":
            header_row_idx = i
            col_headers = [str(c).strip().upper() if c else "" for c in row]
            break

    if header_row_idx is None:
        wb.close()
        raise ValueError("Header baris kolom (NO, KDCAB, ...) tidak ditemukan di Excel!")

    def ci(name: str) -> int:
        for idx, h in enumerate(col_headers):
            if name in h:
                return idx
        raise KeyError(f"Kolom {name} tidak ditemukan di header")

    idx_jenis = ci("JENIS MASALAH")
    idx_kdcab = ci("KDCAB")
    idx_kdtk = ci("KDTK")
    idx_namatk = ci("NAMA TOKO")
    idx_nonota = ci("NONOTA")
    idx_tgl = ci("TGL KIRIM")
    idx_nilaisj = ci("NILAI SJ")
    idx_item = ci("NAMA ITEM") if any("NAMA ITEM" in h for h in col_headers) else None

    today = date.today()
    snapshot_rows = []
    per_jenis_count = defaultdict(int)

    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if not row[0] or not isinstance(row[0], int):
            continue

        kdcab = str(row[idx_kdcab] or "").strip()
        if kdcab not in MASTER_KDCAB:
            continue

        raw_jenis = str(row[idx_jenis] or "").strip().upper()
        jenis = "INTRANSIT" if "INTRANSIT" in raw_jenis else ("SELISIH_BK" if "SELISIH" in raw_jenis else "LAINNYA")
        kdtk = str(row[idx_kdtk] or "").strip()
        nama_toko = str(row[idx_namatk] or "").strip()
        nonota = str(row[idx_nonota] or "").strip()
        tgl_kirim = parse_date_val(row[idx_tgl])
        nilai_sj = float(row[idx_nilaisj] or 0)
        nama_item = str(row[idx_item] or "") if idx_item is not None and row[idx_item] else None

        snapshot_rows.append((
            today, jenis, kdcab, kdtk, nama_toko, nonota, tgl_kirim, nilai_sj, nama_item
        ))
        per_jenis_count[jenis] += 1

    wb.close()

    if not snapshot_rows:
        logger.warning("Tidak ada baris data ditemukan di laporan rekap.")
        return {"tgl_snapshot": str(today), "total_nota": 0, "per_jenis": {}}

    conn = db.get_connection(config.DB_NAME)
    cur = conn.cursor()

    insert_sql = """
        INSERT INTO snapshot_intransit (
            tgl_snapshot, jenis_masalah, kode_cabang, kode_toko, nama_toko,
            docno, tgl_kirim, nilai_sj, nama_item_selisih
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON DUPLICATE KEY UPDATE
            kode_cabang = VALUES(kode_cabang),
            nama_toko = VALUES(nama_toko),
            tgl_kirim = VALUES(tgl_kirim),
            nilai_sj = VALUES(nilai_sj),
            nama_item_selisih = VALUES(nama_item_selisih);
    """
    cur.executemany(insert_sql, snapshot_rows)
    conn.commit()
    cur.close()
    conn.close()

    logger.info(f"✓ Berhasil menyimpan {len(snapshot_rows)} baris snapshot untuk {today}")
    return {
        "status": "success",
        "tgl_snapshot": str(today),
        "total_nota": len(snapshot_rows),
        "per_jenis": dict(per_jenis_count),
    }

def main():
    session = create_authenticated_session()
    content = download_rekap_ringkas(session)
    return parse_and_store_snapshot(content)

if __name__ == "__main__":
    run_job(main)
