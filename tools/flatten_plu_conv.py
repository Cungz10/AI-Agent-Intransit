"""
Fetch detailPriceList dari IIAWeb -> flat-in kolom PLU Conv. -> Excel.

Pakai:  python fetch_pricelist.py 80
        python fetch_pricelist.py 80 81 82     # banyak ID, login cuma sekali

.env:   IIA_BASE_URL, IIA_LOGIN_URL, IIA_USERNAME, IIA_PASSWORD
        (opsional: IIA_RESP_ID, IIA_MENU_ID, IIA_FIELD_USER, IIA_FIELD_PASS)
"""
import os
import re
import sys
from io import StringIO

import pandas as pd
import requests
import urllib3
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

CONV_COL = "PLU Conv."
# Urutan penting: yang lebih panjang dicek dulu ("SB BALI" sebelum "SB")
CHANNELS = ["SB TOMANG", "SB AMBON", "SB MANADO", "SB BALI", "PC BALI",
            "ISJM", "MrB", "IDM", "IGR", "AHO", "MD"]
PLACEHOLDER = {"", "-", "0"}  # artinya "nggak ada kode"
KEEP_CHANNEL = "IDM"          # cuma ambil kanal ini; set None buat ambil semua kanal
EMPTY_CODE = "0"              # nilai kalau PLU nggak punya kode di kanal tsb


# ───────────────────────── config & login ─────────────────────────
def _config():
    """Ambil config IIA dari env pas dipanggil (bukan module-load time)."""
    try:
        return {
            "BASE_URL": os.environ["IIA_BASE_URL"].rstrip("/"),
            "LOGIN_URL": os.environ["IIA_LOGIN_URL"],
            "USERNAME": os.environ["IIA_USERNAME"],
            "PASSWORD": os.environ["IIA_PASSWORD"],
            "RESP_ID": os.environ.get("IIA_RESP_ID", "91"),
            "MENU_ID": os.environ.get("IIA_MENU_ID", "14"),
            "FIELD_USER": os.environ.get("IIA_FIELD_USER", "UserName"),
            "FIELD_PASS": os.environ.get("IIA_FIELD_PASS", "Password"),
        }
    except KeyError as e:
        raise RuntimeError(f"Env var {e} belum diset di .env")


def _extract_csrf(html):
    m = (re.search(r'name="__RequestVerificationToken"[^>]*value="([^"]+)"', html)
         or re.search(r'value="([^"]+)"[^>]*name="__RequestVerificationToken"', html))
    if not m:
        raise RuntimeError("CSRF token nggak ketemu di halaman login.")
    return m.group(1)


def login():
    cfg = _config()
    s = requests.Session()
    s.verify = False
    s.headers.update({
        "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                       "(KHTML, like Gecko) Chrome/154.0.0.0 Safari/537.36"),
        "Accept-Language": "id-ID,id;q=0.9,en-US;q=0.8,en;q=0.7",
    })
    r = s.get(cfg["LOGIN_URL"], timeout=30)
    r.raise_for_status()
    token = _extract_csrf(r.text)

    r = s.post(
        cfg["LOGIN_URL"],
        data={"__RequestVerificationToken": token,
              cfg["FIELD_USER"]: cfg["USERNAME"],
              cfg["FIELD_PASS"]: cfg["PASSWORD"]},
        headers={"Referer": cfg["LOGIN_URL"]},
        timeout=30, allow_redirects=True,
    )
    r.raise_for_status()
    if ".AspNet.ApplicationCookie" not in s.cookies.get_dict():
        raise RuntimeError("Login gagal: cookie sesi nggak kebentuk (cek user/pass atau nama field).")
    return s


def select_responsibility(s):
    """Sekali per session (bikin cookie YummyWebCookies kebentuk, kayak di browser)."""
    cfg = _config()
    base = cfg["BASE_URL"]
    xhr = {"Accept": "text/html, */*; q=0.01", "Referer": f"{base}/",
           "X-Requested-With": "XMLHttpRequest"}
    s.get(f"{base}/MasterPriceLists/Index", params={"_respId": cfg["RESP_ID"]},
          headers=xhr, timeout=30)
    s.get(f"{base}/MasterPriceLists/changePageBreadCrumb", params={"_menuId": cfg["MENU_ID"]},
          headers=xhr, timeout=30)


def fetch_detail_pricelist(s, price_list_id):
    base = _config()["BASE_URL"]
    r = s.get(f"{base}/MasterPriceLists/detailPriceList",
              params={"_priceListId": price_list_id},
              headers={"Accept": "text/html, */*; q=0.01", "Referer": f"{base}/",
                       "X-Requested-With": "XMLHttpRequest"},
              timeout=60, allow_redirects=False)
    if r.status_code in (301, 302, 303):
        raise RuntimeError(f"Kena redirect ke {r.headers.get('Location')} -> sesi nggak valid.")
    r.raise_for_status()
    return r.text


# ───────────────────────── parse tabel ─────────────────────────
def pick_item_table(frames):
    """read_html ngembaliin beberapa tabel (ada pembungkus layout). Ambil tabel item
    yang PLU-nya paling bersih (kode alfanumerik pendek)."""
    best, best_score = None, -1.0
    for df in frames:
        if CONV_COL not in df.columns or "PLU" not in df.columns:
            continue
        plu = df["PLU"].dropna().astype(str).str.strip()
        if plu.empty:
            continue
        score = plu.str.fullmatch(r"[A-Za-z0-9]{4,12}").mean()
        if score > best_score:
            best, best_score = df, score
    if best is None:
        raise RuntimeError("Tabel item (kolom 'PLU' + 'PLU Conv.') nggak ketemu di HTML.")
    best = best[best["PLU"].notna()].copy()
    best["PLU"] = best["PLU"].astype(str).str.strip()
    return best.reset_index(drop=True)


def parse_plu_conv(text):
    """'MD A - 1. IDM 2. IGR .' -> ([(channel, kode), ...], [token_tak_dikenal])"""
    if text is None or (isinstance(text, float) and pd.isna(text)):
        return [], []
    pairs, unknown = [], []
    for tok in re.split(r"\.(?:\s+|$)", str(text).strip()):
        tok = tok.strip()
        if not tok:
            continue
        for ch in CHANNELS:
            if tok == ch or tok.startswith(ch + " "):
                kode = tok[len(ch):].strip()
                if kode.startswith(ch + " "):          # typo sumber: "MD MD NSAR-9343"
                    kode = kode[len(ch):].strip()
                pairs.append((ch, kode))
                break
        else:
            unknown.append(tok)
    return pairs, unknown


def flatten(df, col=CONV_COL):
    rows, unknowns = [], []
    stats = {"dup": 0, "no_code": 0}
    for _, r in df.iterrows():
        pairs, unk = parse_plu_conv(r[col])
        unknowns += [(r["PLU"], u) for u in unk]
        seen, real = set(), []
        for ch, kode in pairs:
            if KEEP_CHANNEL and ch != KEEP_CHANNEL:
                continue
            if kode in PLACEHOLDER:
                continue
            if (ch, kode) in seen:
                stats["dup"] += 1
                continue
            seen.add((ch, kode))
            real.append((ch, kode))
        base = r.drop(labels=[col]).to_dict()
        if not real:                      # PLU tetap muncul 1 baris, kode diisi 0
            stats["no_code"] += 1
            rows.append({**base, "Conv Channel": KEEP_CHANNEL,
                         "Conv Kode": EMPTY_CODE if KEEP_CHANNEL else None})
        for ch, kode in real:
            rows.append({**base, "Conv Channel": ch, "Conv Kode": kode})
    return pd.DataFrame(rows), unknowns, stats


# ───────────────────────── output ─────────────────────────
def save_excel(flat, path):
    flat.to_excel(path, index=False, sheet_name="flat")
    wb = load_workbook(path)
    ws = wb["flat"]
    hdr = PatternFill("solid", start_color="D9E1F2")
    text_cols = {i for i, c in enumerate(ws[1], 1) if c.value in ("PLU", "Conv Kode")}
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, size=10)
        c.fill = hdr
        c.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
            if c.column in text_cols:
                c.number_format = "@"
    for i, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 10), 42)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def process(s, price_list_id):
    html = fetch_detail_pricelist(s, price_list_id)
    with open(f"detail_pricelist_{price_list_id}.html", "w", encoding="utf-8") as f:
        f.write(html)

    try:
        frames = pd.read_html(StringIO(html))
    except ValueError:
        raise RuntimeError(f"[{price_list_id}] Nggak ada <table> di respons, cek HTML mentahnya.")

    items = pick_item_table(frames)
    flat, unknowns, stats = flatten(items)
    out = f"pricelist_{price_list_id}_flat.xlsx"
    save_excel(flat, out)

    print(f"[{price_list_id}] {len(items)} PLU -> {len(flat)} baris | "
          f"duplikat dibuang: {stats['dup']} | PLU tanpa kode (diisi {EMPTY_CODE}): {stats['no_code']} | "
          f"token tak dikenal: {len(unknowns)} -> {out}")
    for plu, u in unknowns:
        print(f"    ? {plu}: {u}")


if __name__ == "__main__":
    ids = [int(a) for a in sys.argv[1:]] or [80]
    session = login()
    print("[+] login OK")
    select_responsibility(session)
    for pid in ids:
        process(session, pid)
